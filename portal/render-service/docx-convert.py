#!/usr/bin/env python
# docx-convert.py — Word(.docx) → Markdown 的「機械轉換」部分（忠實、可重複、可排程）。
# 做：① mammoth 轉 HTML →（markdownify）轉 md：**表格保留成 GFM 表格**、清單用空格縮排（mammoth 的
#        markdown writer 會丟掉表格、且巢狀清單用 Tab→被當程式碼區塊，故改走 HTML→markdownify）；
#     ② 圖片抽成獨立檔（不內嵌 base64，避免 md 爆大）、alt 清空（避免多行 alt 把 ![] 語法弄壞）；
#     ③ EMF/WMF 向量圖（Word 畫的流程圖、Visio/OLE 物件的預覽常是這種）用 Windows System.Drawing
#        轉成 PNG，讓圖能顯示、也讓 AI 看得到；④ 去封面/簽核/版本歷史/目錄樣板；
#     ⑤ 安全網：把 docx 內「全部」內嵌圖（word/media/*）另存到 _media_all/，避免 mammoth 漏放。
# K1-01: inject_para_ids() — 匯入時補齊 w14:paraId（原封不動存回 source.docx）。
# K1-02: convert_import() 改用 python-docx 自走訪 walker 生 doc.md + <!-- el:paraId --> 錨點。
# 不做（交給 AI/skill）：判斷哪些圖是「流程圖」要改畫成 BPMN、哪些是純圖直接內嵌。
# 相依：mammoth、markdownify（pip install mammoth markdownify）; python-docx（walker）。
# 用法：python docx-convert.py "<檔>.docx" "<輸出目錄>"
import sys, os, re, subprocess, zipfile, tempfile, shutil, random as _random
import posixpath
import mammoth
import markdownify

EXT = {'image/png': '.png', 'image/jpeg': '.jpg', 'image/jpg': '.jpg',
       'image/x-emf': '.emf', 'image/x-wmf': '.wmf', 'image/gif': '.gif', 'image/bmp': '.bmp'}

_SENTINEL = 'HANAPBRKSENTINEL2025'

def inject_page_break_sentinels(docx_path):
    """Return (path, is_tmp): temp docx with sentinel paragraphs replacing explicit page-break runs.
    If no page breaks found, returns (docx_path, False) — no temp file created."""
    with zipfile.ZipFile(docx_path, 'r') as z:
        if 'word/document.xml' not in z.namelist():
            return docx_path, False
        doc_xml = z.read('word/document.xml').decode('utf-8')

    if ('<w:br w:type="page"' not in doc_xml and "<w:br w:type='page'" not in doc_xml):
        return docx_path, False

    # Replace each entire <w:p>…</w:p> block that contains a page break
    # with a sentinel-text paragraph.  Any text co-located with the break
    # is intentionally dropped (Word page-break paragraphs are usually empty).
    def _replace_para(m):
        para = m.group(0)
        if re.search(r'''<w:br\s+w:type=['"]page['"]''', para):
            return '<w:p><w:r><w:t>' + _SENTINEL + '</w:t></w:r></w:p>'
        return para

    doc_xml = re.sub(r'<w:p\b.*?</w:p>', _replace_para, doc_xml, flags=re.DOTALL)

    # Write a modified copy to a temp file
    tmp = tempfile.mktemp(suffix='.docx')
    with zipfile.ZipFile(docx_path, 'r') as zin:
        with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.namelist():
                zout.writestr(item, doc_xml.encode('utf-8') if item == 'word/document.xml' else zin.read(item))
    return tmp, True

def emf_to_png(folder):
    """把資料夾內所有 .emf/.wmf 轉成同名 .png（Windows System.Drawing，免裝）。"""
    if os.name != 'nt':
        return
    ad = folder.replace('\\', '/')
    ps = (
        "Add-Type -AssemblyName System.Drawing; "
        "Get-ChildItem '%s/*' -Include *.emf,*.wmf -File | ForEach-Object { try { "
        "$i=[System.Drawing.Image]::FromFile($_.FullName); "
        "$b=New-Object System.Drawing.Bitmap($i.Width,$i.Height); "
        "$g=[System.Drawing.Graphics]::FromImage($b); $g.Clear([System.Drawing.Color]::White); "
        "$g.DrawImage($i,0,0,$i.Width,$i.Height); "
        "$b.Save(($_.FullName -replace '\\.(emf|wmf)$','.png'),[System.Drawing.Imaging.ImageFormat]::Png) "
        "} catch {} }"
    ) % ad
    subprocess.run(['powershell', '-NoProfile', '-Command', ps], capture_output=True)

def convert(docx_path, out_dir):
    name = os.path.splitext(os.path.basename(docx_path))[0]
    # 每個 md 一個素材子資料夾：素材/<md檔名>/。圖片放這、之後 AI 產的 .bpmn / .flow.json 也放這
    # → 主目錄只剩乾淨的 .md，用資料夾當命名空間 → 不同 word 的同名流程圖不會打架。
    assets_rel = '素材/' + name
    assets_dir = os.path.join(out_dir, '素材', name)
    os.makedirs(assets_dir, exist_ok=True)

    counter = {'n': 0}
    images = []
    def handle_image(image):
        counter['n'] += 1
        ext = EXT.get((image.content_type or '').lower(), '.bin')
        fname = 'image%d%s' % (counter['n'], ext)
        with image.open() as f:
            open(os.path.join(assets_dir, fname), 'wb').write(f.read())
        images.append({'file': fname, 'type': image.content_type})
        return {'src': assets_rel + '/' + fname, 'alt': ''}  # alt 留空，避免 Word 的多行圖說把 ![] 弄壞

    # 偵測 Word 明確分頁符：把含 page-break 的段落換成 sentinel 文字段，
    # mammoth 轉完後再把 sentinel 換回 <!-- PAGE_BREAK --> 標記。
    src_docx, is_tmp = inject_page_break_sentinels(docx_path)
    try:
        with open(src_docx, 'rb') as f:
            result = mammoth.convert_to_html(f, convert_image=mammoth.images.img_element(handle_image))
    finally:
        if is_tmp:
            try:
                os.unlink(src_docx)
            except Exception:
                pass
    md = markdownify.markdownify(result.value, heading_style='ATX', bullets='-', strip=['a'])
    md = md.replace('\r\n', '\n').replace('\r', '\n')  # 正規化換行
    md = re.sub(r'<a id="[^"]*"></a>', '', md)  # 清空書籤錨點（markdownify 多半已移除，保險）
    md = md.replace(_SENTINEL, '<!-- PAGE_BREAK -->')  # 還原 Word 明確分頁符
    md = re.sub(r'\n{3,}', '\n\n', md)  # 收掉 markdownify 偶有的連續空行

    # 去掉「封面/簽核/版本歷史/目錄」這類文件控制樣板：偵測到目錄(TOC)時，保留文件標題 +
    # 從目錄後第一個標題開始的正文。沒有 TOC 就保持完整、不亂砍（避免砍到真內容）。
    lines = md.split('\n')
    toc_i = next((i for i, l in enumerate(lines) if re.search(r'Table of Contents|目\s*錄', l)), None)
    if toc_i is not None:
        h_i = next((i for i in range(toc_i, len(lines)) if re.match(r'^#{1,3}\s', lines[i])), None)
        if h_i is not None:
            title_parts = []
            for l in lines[:toc_i]:
                # 封面表格一開始（| …）或簽核樣板就停：標題只取前面那幾行文件名
                if l.lstrip().startswith('|') or re.search(r'DOCUMENT ID|NUMBER OF PAGES|FILENAME|COPY NUMBER|APPROVED BY|ACCEPTED BY|Revision History', l, re.I):
                    break
                t = re.sub(r'[*_#\\]', '', l).strip()  # 去掉 markdown 粗體/標題符號
                if t:
                    title_parts.append(t)
            title = ' '.join(title_parts).strip()
            body = '\n'.join(lines[h_i:])
            md = ('# ' + title + '\n\n' + body) if title else body

    # 修巢狀清單縮排：mammoth 用「Tab」縮排巢狀清單、且偶有「只有項目符號的空項」，
    # Markdown 會把 Tab(=4格) 縮排當成「程式碼區塊」→ 整段變灰、像 code（白底幾乎看不到）。
    # 改成「每層 3 空格、丟掉空項目符號」→ 正常巢狀 <ol>/<ul>。
    norm = []
    for l in md.split('\n'):
        m = re.match(r'^(\t+)(.*)$', l)
        if m:
            stripped = m.group(2)
            if stripped.strip() in ('-', '*', '+'):  # 空項目符號（Word 空層級的產物）→ 丟掉
                continue
            l = '   ' * len(m.group(1)) + stripped     # Tab → 3 空格/層（避免觸發程式碼區塊）
        norm.append(l)
    md = '\n'.join(norm)

    # mammoth 放進 md 的圖：EMF/WMF → PNG，並改 md 引用
    emf_to_png(assets_dir)
    md = md.replace('.emf)', '.png)').replace('.wmf)', '.png)')
    for im in images:
        im['file'] = re.sub(r'\.(emf|wmf)$', '.png', im['file'])

    # 安全網：dump docx 內全部內嵌圖（word/media/*），轉 EMF/WMF→PNG
    media_all = []
    try:
        all_dir = os.path.join(assets_dir, '_media_all')
        z = zipfile.ZipFile(docx_path)
        media = sorted([n for n in z.namelist() if n.startswith('word/media/')])
        if media:
            os.makedirs(all_dir, exist_ok=True)
            for n in media:
                bn = os.path.basename(n)
                open(os.path.join(all_dir, bn), 'wb').write(z.read(n))
            emf_to_png(all_dir)
            media_all = sorted(os.listdir(all_dir))
    except Exception:
        pass

    md_path = os.path.join(out_dir, name + '.md')
    open(md_path, 'w', encoding='utf-8').write(md)

    # 給 AI 接手的清單
    print('MD_FILE: ' + md_path)
    print('ASSETS_DIR: ' + assets_dir + '  （此 md 的素材夾；圖片在這，之後 .flow.json / .bpmn 也放這）')
    print('IMAGES_IN_MD (mammoth 已放進 md；逐張看圖：流程圖→寫 .flow.json→產 .bpmn 內嵌；純圖→保留 ![](%s/<檔>)):' % assets_rel)
    for im in images:
        print('  - %s/%s  [%s]' % (assets_rel, im['file'], im['type']))
    pngs_all = [b for b in media_all if not b.lower().endswith(('.emf', '.wmf'))]
    print('ALL_EMBEDDED_IMAGES (%s/_media_all/，docx 內全部內嵌圖，含 mammoth 可能漏放的；' % assets_rel)
    print('  若這裡有「流程圖」但上面 IMAGES_IN_MD 沒有，代表 mammoth 漏放 → 請對照原文位置自行插入 ```bpmn 或 ![]):')
    for b in pngs_all:
        print('  - %s/_media_all/%s' % (assets_rel, b))
    if not images and not pngs_all:
        print('  （無內嵌圖；若 Word 圖是 DrawingML 畫的形狀，抽不到，請指揮官貼截圖）')
    for w in result.messages[:5]:
        if 'warning' in str(w).lower():
            print('WARN: ' + str(w)[:160])

# ── Front-matter utilities (DOC-I1-03) ───────────────────────────────────────
# Simple YAML-subset front-matter: scalar strings + arrays written as [a, b, c].
# front-matter is source of truth; registry is a cache.

def parse_frontmatter(content):
    """Return (meta: dict, body: str). meta is empty dict if no front-matter found."""
    import re as _re
    m = _re.match(r'^---\r?\n(.*?)\r?\n---\r?\n?', content, _re.DOTALL)
    if not m:
        return {}, content
    meta = {}
    for line in m.group(1).split('\n'):
        kv = _re.match(r'^(\w+):\s*(.*)', line)
        if not kv:
            continue
        val = kv.group(2).strip()
        if val.startswith('[') and val.endswith(']'):
            val = [v.strip() for v in val[1:-1].split(',') if v.strip()]
        meta[kv.group(1)] = val
    return meta, content[m.end():]

def stringify_frontmatter(meta):
    lines = []
    for k, v in meta.items():
        if isinstance(v, list):
            lines.append('%s: [%s]' % (k, ', '.join(v)))
        else:
            lines.append('%s: %s' % (k, v))
    return '---\n' + '\n'.join(lines) + '\n---\n'

def write_frontmatter(md_path, meta):
    """Merge meta into the file's existing front-matter (new keys win), preserve body."""
    body = ''
    existing = {}
    if os.path.exists(md_path):
        with open(md_path, encoding='utf-8') as f:
            content = f.read()
        existing, body = parse_frontmatter(content)
    merged = {**existing, **meta}
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write(stringify_frontmatter(merged) + body)


def create_pending_bpmn(bpmn_path, image_ref):
    """Create a placeholder .bpmn file for flowcharts that need manual/Hana reconstruction."""
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<bpmn:definitions xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL"\n'
        '  xmlns:bpmndi="http://www.omg.org/spec/BPMN/20100524/DI"\n'
        '  targetNamespace="http://harness.ai">\n'
        '  <!-- 待補（人工/Hana）：此流程圖來自 Visio/OLE 向量圖，需根據原始圖片重建。\n'
        '       原始圖片：%s -->\n'
        '  <bpmn:process id="Process_pending" isExecutable="false">\n'
        '    <bpmn:startEvent id="start_pending"\n'
        '      name="待補&#xa;請根據 %s&#xa;重建此流程圖" />\n'
        '  </bpmn:process>\n'
        '</bpmn:definitions>\n'
    ) % (image_ref, image_ref)
    with open(bpmn_path, 'w', encoding='utf-8') as f:
        f.write(xml)


# ── K1: paraId injection + python-docx walker ─────────────────────────────────

_NS_W14 = 'http://schemas.microsoft.com/office/word/2010/wordml'
_NS_W   = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
_NS_WP  = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
_NS_A   = 'http://schemas.openxmlformats.org/drawingml/2006/main'
_NS_R   = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
_W14 = '{%s}' % _NS_W14
_W   = '{%s}' % _NS_W
_WP  = '{%s}' % _NS_WP
_A   = '{%s}' % _NS_A
_R   = '{%s}' % _NS_R

_HEADING_PREFIXES = {
    # Style display names (with space)
    'heading 1': '#',   'heading 2': '##',  'heading 3': '###',
    'heading 4': '####','heading 5': '#####','heading 6': '######',
    # Style IDs in XML pStyle val (no space; python-docx default)
    'heading1':  '#',   'heading2':  '##',  'heading3':  '###',
    'heading4':  '####','heading5':  '#####','heading6':  '######',
    # Chinese Word heading styles
    '標題 1': '#', '標題 2': '##', '標題 3': '###',
    '標題1':  '#', '標題2':  '##', '標題3':  '###',
    'title': '#',
}

# ── K1-01 ─────────────────────────────────────────────────────────────────────

def inject_para_ids(docx_path):
    """K1-01: Add w14:paraId to every <w:p> that lacks one. Modifies source.docx in-place.
    Returns count of newly injected IDs. Idempotent: re-running is safe (existing IDs untouched)."""
    with zipfile.ZipFile(docx_path, 'r') as z:
        if 'word/document.xml' not in z.namelist():
            return 0
        doc_xml = z.read('word/document.xml').decode('utf-8')
        all_items = [(n, z.read(n)) for n in z.namelist()]

    existing = set(re.findall(r'w14:paraId="([0-9A-Fa-f]{8})"', doc_xml, re.I))
    injected = [0]

    def _add_pid(m):
        tag = m.group(0)
        if 'w14:paraId' in tag:
            return tag
        while True:
            pid = '%08X' % _random.randint(0, 0xFFFFFFFF)
            if pid not in existing:
                existing.add(pid)
                break
        injected[0] += 1
        if tag.endswith('/>'):
            return tag[:-2] + ' w14:paraId="%s"/>' % pid
        return tag[:-1] + ' w14:paraId="%s">' % pid

    # Match only <w:p> opening tags (word boundary prevents matching <w:pPr> etc.)
    doc_xml = re.sub(r'<w:p\b[^>]*/?>', _add_pid, doc_xml)

    # Ensure w14 namespace is declared on document root
    if 'xmlns:w14=' not in doc_xml and injected[0] > 0:
        doc_xml = re.sub(
            r'(<w:document\b[^>]*?)((?:\s*/?>|\s+xmlns:))',
            r'\1 xmlns:w14="%s"\2' % _NS_W14,
            doc_xml, count=1
        )

    with zipfile.ZipFile(docx_path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for name, data in all_items:
            zout.writestr(name,
                doc_xml.encode('utf-8') if name == 'word/document.xml' else data)

    return injected[0]

# ── K1-02 helpers ─────────────────────────────────────────────────────────────

_EXT_BY_EXT = {'.png': '.png', '.jpg': '.jpg', '.jpeg': '.jpg',
               '.gif': '.gif', '.bmp': '.bmp', '.emf': '.emf', '.wmf': '.wmf',
               '.svg': '.svg', '.tif': '.tif', '.tiff': '.tif'}

def _build_img_map(docx_path, assets_dir):
    """Extract all images from docx zip to assets_dir.
    Returns ({rId: 'assets/imageN.ext'}, set_of_emf_rids).
    EMF/WMF are converted to PNG in-place; map already reflects the .png filename."""
    img_map = {}
    emf_rids = set()

    with zipfile.ZipFile(docx_path, 'r') as z:
        names = set(z.namelist())
        rels_data = ''
        rels_path = 'word/_rels/document.xml.rels'
        if rels_path in names:
            rels_data = z.read(rels_path).decode('utf-8', errors='replace')

        # Parse each Relationship element regardless of attribute order
        img_rels = []
        for m in re.finditer(r'<Relationship\b([^>]+?)/?\s*>', rels_data):
            attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
            rtype = attrs.get('Type', '')
            if '/image' not in rtype and 'image' not in rtype.lower():
                continue
            rid = attrs.get('Id', '')
            target = attrs.get('Target', '')
            if rid and target:
                img_rels.append((rid, target))

        counter = [0]
        for rid, target in img_rels:
            # Resolve target relative to word/
            zip_path = posixpath.normpath('word/' + target)
            if zip_path not in names:
                zip_path = target.lstrip('/')
            if zip_path not in names:
                continue

            ext = os.path.splitext(target)[1].lower()
            out_ext = _EXT_BY_EXT.get(ext, '.bin')
            counter[0] += 1
            fname = 'image%d%s' % (counter[0], out_ext)

            with open(os.path.join(assets_dir, fname), 'wb') as f:
                f.write(z.read(zip_path))

            img_map[rid] = 'assets/' + fname
            if ext in ('.emf', '.wmf'):
                emf_rids.add(rid)

    # EMF/WMF → PNG
    emf_to_png(assets_dir)
    img_map = {rid: re.sub(r'\.(emf|wmf)$', '.png', path, flags=re.I)
               for rid, path in img_map.items()}

    return img_map, emf_rids


# ── flowmap sidecar (§13.8-5) ────────────────────────────────────────────────
# 匯入時把「流程圖(EMF)→bpmn placeholder」的那句 re.sub 連同圖後的 docPr/el 錨點註解一起吃掉，
# 導致「擬真畫面某流程圖 ↔ 哪個 assets/*.bpmn」對應遺失（擋住 P3-b 合成、乙 bpmn→svg 換回）。
# 對策：額外寫 flowmap.json sidecar（與 assets/ 同層），記每張被轉成 bpmn 的流程圖的對應。
# 關鍵：真實客戶檔的流程圖多為 Visio/OLE 物件（<w:object>/<v:imagedata>，**沒有 wp:docPr**），
# 少數才是 DrawingML(<a:blip>+docPr)；所以可靠的鍵是「文件順序 index」(flowIndex)，docPrId 僅在
# DrawingML 時有值。paraId＝該圖所在段落錨點（docx-preview 的 data-el），為額外定位提示。

def _image_rel_targets(docx_path):
    """{rId: target} for image relationships (rels-file order irrelevant). Read-only, no extraction."""
    out = {}
    with zipfile.ZipFile(docx_path, 'r') as z:
        names = set(z.namelist())
        rels_path = 'word/_rels/document.xml.rels'
        if rels_path not in names:
            return out
        rels = z.read(rels_path).decode('utf-8', errors='replace')
    for m in re.finditer(r'<Relationship\b([^>]+?)/?\s*>', rels):
        attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        rtype = attrs.get('Type', '')
        if '/image' not in rtype and 'image' not in rtype.lower():
            continue
        rid = attrs.get('Id', ''); target = attrs.get('Target', '')
        if rid and target:
            out[rid] = target
    return out


def _media_path(target):
    """rels Target → 'word/media/imageN.emf' (targets are word/-relative unless absolute)."""
    t = target.lstrip('/')
    return t if t.startswith('word/') else 'word/' + t


def extract_flow_anchors(docx_path):
    """Doc-order list of the EMF/WMF vector images that become BPMN placeholders (Word flowcharts:
    Visio/OLE preview or DrawingML EMF). Container-agnostic — catches both DrawingML `a:blip` and
    VML/OLE `v:imagedata`. Purely read-only (no file extraction, no conversion).
    Each entry: {rid, docPrId(str|None), paraId(str|None), sourceMedia('word/media/…' | None)}.
    docPrId is None for OLE objects (they carry no wp:docPr); the reliable ordinal key is the list
    position (flowIndex). paraId = containing <w:p> anchor (docx-preview `data-el`)."""
    from docx import Document
    rt  = _image_rel_targets(docx_path)
    emf = {rid for rid, t in rt.items() if os.path.splitext(t)[1].lower() in ('.emf', '.wmf')}
    doc = Document(docx_path)
    out = []
    for p in doc.element.body.iter(_W + 'p'):
        para_id = p.get(_W14 + 'paraId')
        for node in p.iter():
            tag = node.tag.split('}')[-1]
            rid = docpr = None
            if tag == 'drawing':
                blip = next(node.iter(_A + 'blip'), None)
                rid = blip.get(_R + 'embed') if blip is not None else None
                dp = next(node.iter(_WP + 'docPr'), None)
                docpr = dp.get('id') if dp is not None else None
            elif tag in ('object', 'pict'):
                for sub in node.iter():
                    if sub.tag.split('}')[-1] == 'imagedata':
                        rid = sub.get(_R + 'id'); break
            if rid and rid in emf:
                out.append({'rid': rid, 'docPrId': docpr, 'paraId': para_id,
                            'sourceMedia': _media_path(rt[rid])})
    # One flowchart can surface the same rid twice within a paragraph (e.g. object wrapping a
    # drawing preview) → collapse consecutive duplicates of (rid, paraId).
    dedup = []
    for a in out:
        if dedup and dedup[-1]['rid'] == a['rid'] and dedup[-1]['paraId'] == a['paraId']:
            continue
        dedup.append(a)
    return dedup


def _bpmn_ref_for_image(img_rel):
    """assets/imageN.png → assets/imageN_pending.bpmn (import creates the _pending placeholder)."""
    return os.path.splitext(img_rel)[0] + '_pending.bpmn'


def _image_for_bpmn_ref(ref):
    """assets/imageN[_pending].bpmn → assets/imageN.png (best-effort source-image name)."""
    base = re.sub(r'_pending\.bpmn$', '.bpmn', ref, flags=re.I)
    return re.sub(r'\.bpmn$', '.png', base, flags=re.I)


def write_flowmap(doc_dir, docx_path, img_map, source_rel='', generated_by='import'):
    """Write flowmap.json: each BPMN placeholder ↔ its source flowchart image, in document order.
    Additive — never touches doc.md / assets / bpmn. Returns entry count (0 → no sidecar written)."""
    import json as _json
    anchors = extract_flow_anchors(docx_path)
    entries = []
    for i, a in enumerate(anchors):
        img_rel = img_map.get(a['rid'])
        if not img_rel:
            continue
        entries.append({
            'flowIndex':   i,
            'docPrId':     a['docPrId'],
            'paraId':      a['paraId'],
            'sourceMedia': a['sourceMedia'],
            'sourceImage': img_rel,
            'bpmnRef':     _bpmn_ref_for_image(img_rel),
        })
    if not entries:
        return 0
    fm = {'version': 1, 'source': source_rel or os.path.basename(docx_path),
          'generatedBy': generated_by, 'entries': entries}
    with open(os.path.join(doc_dir, 'flowmap.json'), 'w', encoding='utf-8') as f:
        _json.dump(fm, f, ensure_ascii=False, indent=2)
    return len(entries)


def backfill_flowmap(workspace_root):
    """One-shot backfill of flowmap.json for already-imported docs. READ-ONLY except for writing the
    new sidecar: never re-imports, re-converts, re-generates bpmn, or touches doc.md/assets/source.
    Per doc with ```bpmn fences: re-extract the EMF flowchart images (document order) from the
    ORIGINAL docx (registry sourcePath = workspace + sourcePath) and pair 1:1 by document order with
    the fences. Honest guardrail — do not force a fit: if the original is missing, unparseable, or its
    flowchart-image count != fence count (order can't be trusted; likely hand-edited or heterogeneous
    importer), skip and list the doc under needsReview. Prints + returns a JSON report."""
    import json as _json
    docs_dir = os.path.join(workspace_root, '.documents')
    idx = os.path.join(docs_dir, 'index.json')
    report = {'workspace': workspace_root, 'written': [], 'needsReview': []}
    if not os.path.exists(idx):
        report['error'] = 'no registry at ' + idx
        print(_json.dumps(report, ensure_ascii=False, indent=2))
        return report
    reg = _json.load(open(idx, encoding='utf-8'))
    for d in reg.get('documents', []):
        if d.get('type') != 'doc':
            continue
        doc_dir = os.path.join(docs_dir, d['id'])
        md_path = os.path.join(doc_dir, 'doc.md')
        if not os.path.exists(md_path):
            continue
        md = open(md_path, encoding='utf-8').read()
        fences = re.findall(r'```bpmn[ \t]*\r?\n(.*?)\r?\n```', md, re.S)
        frefs = [f.strip().splitlines()[0].strip() for f in fences if f.strip()]
        frefs = [r for r in frefs if r and not r.startswith('<')]   # skip inline-XML fences
        if not frefs:
            continue
        name = d.get('sourcePath', d['id'])
        src_rel = d.get('sourcePath') or ''
        src = os.path.join(workspace_root, src_rel.replace('/', os.sep))
        if not src_rel or not os.path.exists(src):
            report['needsReview'].append({'id': d['id'], 'source': name,
                'reason': '原始 docx 缺失', 'fences': len(frefs)})
            continue
        try:
            anchors = extract_flow_anchors(src)
        except Exception as e:
            report['needsReview'].append({'id': d['id'], 'source': name,
                'reason': '原始 docx 解析失敗: %s' % e, 'fences': len(frefs)})
            continue
        if len(anchors) != len(frefs):
            report['needsReview'].append({'id': d['id'], 'source': name,
                'reason': '流程圖(EMF)數 %d ≠ bpmn fence 數 %d，順序配不齊' % (len(anchors), len(frefs)),
                'fences': len(frefs)})
            continue
        entries = [{
            'flowIndex':   i,
            'docPrId':     a['docPrId'],
            'paraId':      a['paraId'],
            'sourceMedia': a['sourceMedia'],
            'sourceImage': _image_for_bpmn_ref(ref),
            'bpmnRef':     ref,
        } for i, (a, ref) in enumerate(zip(anchors, frefs))]
        fm = {'version': 1, 'source': name, 'generatedBy': 'backfill', 'entries': entries}
        with open(os.path.join(doc_dir, 'flowmap.json'), 'w', encoding='utf-8') as f:
            _json.dump(fm, f, ensure_ascii=False, indent=2)
        report['written'].append({'id': d['id'], 'source': name, 'entries': len(entries)})
    report['summary'] = {'written': len(report['written']),
                         'needsReview': len(report['needsReview'])}
    print(_json.dumps(report, ensure_ascii=False, indent=2))
    return report


def _run_text(r_elem):
    """Extract text from a <w:r> with bold/italic markdown decoration."""
    rPr = r_elem.find(_W + 'rPr')
    bold = italic = False
    if rPr is not None:
        b = rPr.find(_W + 'b')
        i = rPr.find(_W + 'i')
        if b is not None and b.get(_W + 'val', '1') not in ('0', 'false', 'off'):
            bold = True
        if i is not None and i.get(_W + 'val', '1') not in ('0', 'false', 'off'):
            italic = True
    t = r_elem.find(_W + 't')
    if t is None or not t.text:
        return ''
    text = t.text
    if bold and italic: return '***' + text + '***'
    if bold:            return '**'  + text + '**'
    if italic:          return '*'   + text + '*'
    return text


def _build_list_styles(styles_elem):
    """Return set of style IDs (lowercase) that declare list/numbering in their pPr.
    styles_elem: the <w:styles> lxml element (doc.part.styles.element).
    Used as fallback when a paragraph's own pPr has no numPr but the style does."""
    result = set()
    for style in styles_elem.findall(_W + 'style'):
        sid = (style.get(_W + 'styleId') or '').lower()
        if not sid:
            continue
        pPr = style.find(_W + 'pPr')
        if pPr is None:
            continue
        numPr = pPr.find(_W + 'numPr')
        if numPr is None:
            continue
        ni = numPr.find(_W + 'numId')
        if ni is not None and ni.get(_W + 'val', '0') != '0':
            result.add(sid)
    return result


def _build_heading_styles(styles_elem):
    """Return {styleId (lowercase): markdown heading level 1..9} for paragraph styles that carry an
    <w:outlineLvl> (directly or inherited via <w:basedOn>). outlineLvl N → level N+1.

    WHY (2026-07-03 regression fix): heading styles can have arbitrary/numeric styleIds — this real
    client docx uses styleId "1"/"2" whose NAMES are "heading 1"/"heading 2". _para_to_md_block matches
    _HEADING_PREFIXES by the raw styleId ("1"/"2"), which is neither "heading1" nor "heading 1" → the
    heading was missed and fell through to bullet/text (doc.md lost all `#`). `outlineLvl` is the
    canonical, name-independent heading marker (what the old mammoth path used), so resolve level from
    it. TOC styles (e.g. "toc 1") have no outlineLvl → correctly NOT treated as headings."""
    raw, based = {}, {}
    for style in styles_elem.findall(_W + 'style'):
        if (style.get(_W + 'type') or '') != 'paragraph':
            continue
        sid = (style.get(_W + 'styleId') or '').lower()
        if not sid:
            continue
        lvl = None
        pPr = style.find(_W + 'pPr')
        if pPr is not None:
            ol = pPr.find(_W + 'outlineLvl')
            if ol is not None:
                try:
                    lvl = int(ol.get(_W + 'val', ''))
                except (TypeError, ValueError):
                    lvl = None
        raw[sid] = lvl
        bo = style.find(_W + 'basedOn')
        if bo is not None:
            based[sid] = (bo.get(_W + 'val') or '').lower()
    result = {}
    for sid in raw:
        cur, seen, lvl = sid, set(), None
        while cur and cur not in seen:      # walk basedOn chain (cycle-safe)
            seen.add(cur)
            if raw.get(cur) is not None:
                lvl = raw[cur]
                break
            cur = based.get(cur)
        if lvl is not None and 0 <= lvl <= 8:
            result[sid] = lvl + 1
    return result


def _para_to_md_block(p_elem, img_map, list_styles=None, heading_styles=None):
    """Convert a <w:p> lxml element to (md_text, anchor_id).
    anchor_id: w14:paraId for text paras; wp:docPr id for image-only paras.
    list_styles: set of style IDs (lowercase) that have list formatting in the style definition."""
    para_id = p_elem.get(_W14 + 'paraId')

    # Page break detection
    for br in p_elem.iter(_W + 'br'):
        if br.get(_W + 'type') == 'page':
            return '<!-- PAGE_BREAK -->', para_id

    # Paragraph properties
    pPr = p_elem.find(_W + 'pPr')
    style_val = ''
    num_id = None
    ilvl = 0
    if pPr is not None:
        ps = pPr.find(_W + 'pStyle')
        if ps is not None:
            style_val = (ps.get(_W + 'val') or '').lower()
        numPr = pPr.find(_W + 'numPr')
        if numPr is not None:
            ni = numPr.find(_W + 'numId')
            il = numPr.find(_W + 'ilvl')
            if ni is not None: num_id = ni.get(_W + 'val')
            if il is not None:
                try: ilvl = int(il.get(_W + 'val', '0'))
                except: ilvl = 0

    # Fallback: detect list via style definition (numPr in style, not paragraph pPr)
    if not num_id and list_styles and style_val in list_styles:
        num_id = 'style'  # sentinel: has list formatting via style

    text_parts = []
    img_refs = []   # (img_path, docpr_id) for image-only paragraphs

    for r in p_elem.findall(_W + 'r'):
        drawing = r.find(_W + 'drawing')
        if drawing is not None:
            docpr_id = None
            for dp in drawing.iter(_WP + 'docPr'):
                docpr_id = dp.get('id')
                break
            embed_id = None
            for blip in drawing.iter(_A + 'blip'):
                embed_id = blip.get(_R + 'embed')
                break
            if embed_id and embed_id in img_map:
                img_refs.append((img_map[embed_id], docpr_id))
            continue
        text_parts.append(_run_text(r))

    text = ''.join(text_parts).strip()

    # Image-only paragraph: use docPr id as anchor
    if img_refs and not text:
        img_path, docpr_id = img_refs[0]
        anchor = docpr_id or para_id
        return '![](%s)' % img_path, anchor

    # Build markdown text.
    # Heading level: prefer the style's outlineLvl (robust to numeric/custom styleIds like "1"/"2"
    # whose display name is "heading 1"/"heading 2"), fall back to name/id-based _HEADING_PREFIXES.
    _hlvl = (heading_styles or {}).get(style_val)
    heading_pfx = ('#' * _hlvl) if _hlvl else _HEADING_PREFIXES.get(style_val, '')
    if heading_pfx:
        md = (heading_pfx + ' ' + text) if text else ''
    elif num_id and num_id != '0':
        md = ('   ' * ilvl) + '- ' + text
    else:
        md = text

    return md, para_id


def _cell_text(tc_elem):
    """Collect plain text from all paragraphs in a table cell."""
    parts = []
    for p in tc_elem.findall(_W + 'p'):
        cell_parts = []
        for r in p.findall(_W + 'r'):
            t = r.find(_W + 't')
            if t is not None and t.text:
                cell_parts.append(t.text)
        txt = ''.join(cell_parts).strip()
        if txt:
            parts.append(txt)
    return ' '.join(parts)


def _table_to_md_block(tbl_elem):
    """Convert a <w:tbl> element to (md_text, anchor_id).
    anchor_id: w14:paraId of first cell's first paragraph."""
    rows = []
    first_anchor = None

    for tr in tbl_elem.findall(_W + 'tr'):
        cells = []
        for tc in tr.findall(_W + 'tc'):
            if first_anchor is None:
                for fp in tc.findall(_W + 'p'):
                    first_anchor = fp.get(_W14 + 'paraId')
                    break
            cells.append(_cell_text(tc).replace('|', '\\|'))
        if cells:
            rows.append(cells)

    if not rows:
        return '', first_anchor

    n_cols = max(len(r) for r in rows)
    rows = [r + [''] * (n_cols - len(r)) for r in rows]
    lines = ['| ' + ' | '.join(rows[0]) + ' |',
             '|' + ' --- |' * n_cols]
    for row in rows[1:]:
        lines.append('| ' + ' | '.join(row) + ' |')
    return '\n'.join(lines), first_anchor


def convert_import(docx_path, doc_dir, source_rel=''):
    """Import mode (DOC-I2-01/04, K1-01/02): convert docx into .documents/<uuid>/ structure.
    K1-01: inject_para_ids() fills missing w14:paraId in source.docx.
    K1-02: python-docx walker replaces mammoth, emits doc.md + <!-- el:paraId --> anchors.
    Prints MD_FILE / ASSETS_DIR on stdout for server parsing."""
    import datetime
    from docx import Document

    assets_dir = os.path.join(doc_dir, 'assets')
    os.makedirs(assets_dir, exist_ok=True)

    # ── K1-01: inject paraIds into source.docx in-place ──────────────────────
    n_injected = inject_para_ids(docx_path)
    if n_injected:
        print('PARA_IDS_INJECTED: %d' % n_injected)

    # ── K1-02: extract images + walk body ────────────────────────────────────
    img_map, emf_rids = _build_img_map(docx_path, assets_dir)

    doc = Document(docx_path)
    body = doc.element.body
    try:
        list_styles = _build_list_styles(doc.part.styles.element)
        heading_styles = _build_heading_styles(doc.part.styles.element)
    except Exception:
        list_styles = set()
        heading_styles = {}

    blocks = []   # list of (md_text, anchor_id)
    for child in body:
        local = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if local == 'p':
            md, anchor = _para_to_md_block(child, img_map, list_styles, heading_styles)
            blocks.append((md, anchor))
        elif local == 'tbl':
            md, anchor = _table_to_md_block(child)
            blocks.append((md, anchor))
        # sectPr, sdt, etc. → skip

    # Build markdown with anchor comments
    out_lines = []
    for md, anchor in blocks:
        if not md or not md.strip():
            out_lines.append('')
            continue
        if md.startswith('<!-- '):
            # PAGE_BREAK or bare comment: no extra anchor
            out_lines.append(md)
        elif '\n' in md:
            # Multi-line block (table): anchor on its own line after the block
            out_lines.append(md)
            if anchor:
                out_lines.append('<!-- el:%s -->' % anchor)
        else:
            anc = ('  <!-- el:%s -->' % anchor) if anchor else ''
            out_lines.append(md + anc)
        out_lines.append('')

    md = '\n'.join(out_lines)
    md = md.replace('\r\n', '\n').replace('\r', '\n')
    md = re.sub(r'\n{3,}', '\n\n', md)

    # TOC / cover page stripping (same logic as convert())
    lines = md.split('\n')
    toc_i = next((i for i, l in enumerate(lines)
                  if re.search(r'Table of Contents|目\s*錄', re.sub(r'<!--[^>]*-->', '', l))), None)
    if toc_i is not None:
        h_i = next((i for i in range(toc_i, len(lines)) if re.match(r'^#{1,3}\s', lines[i])), None)
        if h_i is not None:
            title_parts = []
            for l in lines[:toc_i]:
                l_plain = re.sub(r'\s*<!--[^>]*-->', '', l).strip()
                if l_plain.lstrip().startswith('|') or re.search(
                        r'DOCUMENT ID|NUMBER OF PAGES|FILENAME|COPY NUMBER|APPROVED BY|ACCEPTED BY|Revision History',
                        l_plain, re.I):
                    break
                t = re.sub(r'[*_#\\]', '', l_plain).strip()
                if t:
                    title_parts.append(t)
            title_from_doc = ' '.join(title_parts).strip() or None
            body_md = '\n'.join(lines[h_i:])
            md = ('# ' + title_from_doc + '\n\n' + body_md) if title_from_doc else body_md

    # ── I2-04: EMF/WMF-derived images → bpmn "待補" placeholder ─────────────
    for rid in emf_rids:
        if rid not in img_map:
            continue
        img_path = img_map[rid]   # already converted to .png by _build_img_map
        bpmn_fname = os.path.splitext(os.path.basename(img_path))[0] + '_pending.bpmn'
        bpmn_path = os.path.join(assets_dir, bpmn_fname)
        bpmn_ref = 'assets/' + bpmn_fname
        create_pending_bpmn(bpmn_path, img_path)
        # 路徑寫在 fence「內文」，渲染器讀 code.textContent 拿得到（見 resolveBpmnPath）
        placeholder = (
            '\n\n> ⚠️ **待補（人工/Hana）**：此圖為 Visio/OLE 向量流程圖，'
            '請根據 `%s` 重建 BPMN。\n\n'
            '```bpmn\n%s\n```\n\n'
        ) % (img_path, bpmn_ref)
        # Replace the ![]() line (possibly followed by anchor comment)
        md = re.sub(r'!\[\]\(' + re.escape(img_path) + r'\)(\s*<!--[^>]*-->)?', placeholder, md)

    # Safety net: dump all embedded images from docx
    try:
        all_dir = os.path.join(assets_dir, '_media_all')
        z = zipfile.ZipFile(docx_path)
        media = sorted([n for n in z.namelist() if n.startswith('word/media/')])
        if media:
            os.makedirs(all_dir, exist_ok=True)
            for n in media:
                bn = os.path.basename(n)
                open(os.path.join(all_dir, bn), 'wb').write(z.read(n))
            emf_to_png(all_dir)
    except Exception:
        pass

    now = datetime.datetime.utcnow().strftime('%Y-%m-%dT%H:%M:%SZ')
    fm = {
        'type': 'doc',
        'source': source_rel or os.path.basename(docx_path),
        'convertedAt': now,
        'export': ['docx'],
    }

    md_path = os.path.join(doc_dir, 'doc.md')
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write(stringify_frontmatter(fm) + md)

    # ── flowmap sidecar (§13.8-5): 流程圖(EMF)→bpmn placeholder 對應，doc 順序（additive）──────
    try:
        n_flow = write_flowmap(doc_dir, docx_path, img_map, source_rel, 'import')
        if n_flow:
            print('FLOWMAP_WRITTEN: %d' % n_flow)
    except Exception as e:
        print('FLOWMAP_WARN: %s' % e)   # 只多產一檔，失敗絕不影響 md 匯入結果

    print('MD_FILE: ' + md_path)
    print('ASSETS_DIR: ' + assets_dir)


if __name__ == '__main__':
    if '--import' in sys.argv:
        # --import <docx_path> <doc_dir> [--source <rel_path>]
        args = sys.argv[1:]
        ii = args.index('--import')
        rem = args[ii + 1:]
        if len(rem) < 2:
            print('用法：python docx-convert.py --import "<docx>" "<doc_dir>" [--source "<rel>"]')
            sys.exit(1)
        docx_p = rem[0]
        doc_d  = rem[1]
        src_rel = ''
        if '--source' in rem:
            si = rem.index('--source')
            src_rel = rem[si + 1] if si + 1 < len(rem) else ''
        convert_import(docx_p, doc_d, src_rel)
    elif '--backfill-flowmap' in sys.argv:
        # --backfill-flowmap <workspace_root>  (one-shot; read-only except new flowmap.json)
        args = sys.argv[1:]
        bi = args.index('--backfill-flowmap')
        ws = args[bi + 1] if bi + 1 < len(args) else os.getcwd()
        backfill_flowmap(ws)
    elif '--flow-anchors' in sys.argv:
        # --flow-anchors <docx_path>  → print JSON [{rid,docPrId,paraId,sourceMedia}] in doc order.
        # READ-ONLY. Feeds the fidelity flow overlay's paraId-based anchoring (imageN → paraId).
        import json as _json
        args = sys.argv[1:]
        fi = args.index('--flow-anchors')
        docx_p = args[fi + 1] if fi + 1 < len(args) else ''
        try:
            anchors = extract_flow_anchors(docx_p)
        except Exception as _e:
            anchors = []
        print(_json.dumps(anchors, ensure_ascii=False))
    elif len(sys.argv) >= 3:
        convert(sys.argv[1], sys.argv[2])
    else:
        print('用法：python docx-convert.py "<檔>.docx" "<輸出目錄>"')
        print('      python docx-convert.py --import "<檔>.docx" "<doc_dir>" [--source "<rel>"]')
        sys.exit(1)
